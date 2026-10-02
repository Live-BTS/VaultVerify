#!/usr/bin/env python3
"""Build the MyZipVault Skills Checklist Import Template (.xlsx).

Sheets:
  Instructions — rules, rating scale, question types, format examples
  Skills Data  — 7 canonical columns + dropdown validation + starter content
                 (Zipvault spec example rows + VaultVerify Med-Surg & ICU sets)
"""
import sys, os

XLSX_SKILL_DIR = "/home/z/my-project/skills/xlsx"
for sub in [XLSX_SKILL_DIR, os.path.join(XLSX_SKILL_DIR, "templates")]:
    if sub not in sys.path:
        sys.path.insert(0, sub)

import templates.base as base
from templates.base import (
    setup_sheet, style_header_row, style_data_row,
    font_title, font_subheader, font_body, font_caption,
    align_text, align_header, auto_fit_columns,
)
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

# ── VaultVerify brand palette (from the logo gradient) ──────────
# English-only deliverable for US clinical audiences -> Calibri (Excel-native),
# bold headers allowed. Single-source override, not scattered hardcoding.
base.FONT_NAME = "Calibri"
base.HEADER_BOLD = True

VAULT_PALETTE = {
    "PRIMARY": "03363D",        # Deep Vault Teal
    "PRIMARY_LIGHT": "D8E8DF",  # light mint
    "SECONDARY": "D8E8DF",
    "ACCENT_POSITIVE": "5F940D",  # Verify Bright Green (readable on white)
    "ACCENT_NEGATIVE": "C0392B",
    "ACCENT_WARNING": "D4820A",
    "NEUTRAL_900": "37352F", "NEUTRAL_600": "8C8A84", "NEUTRAL_200": "E9E9E8",
    "NEUTRAL_100": "F7F7F5", "NEUTRAL_50": "FAFAF9", "NEUTRAL_0": "FFFFFF",
    "HEADER_TEXT": "FFFFFF",
    "CHART_COLORS": ["03363D", "5F940D", "D4820A", "C0392B", "8C8A84"],
    "CF_POSITIVE_BG": "E8F5E9", "CF_NEGATIVE_BG": "FDEDEC", "CF_WARNING_BG": "FEF9E7",
}
base._apply(VAULT_PALETTE, "vaultverify")

OUT = "/home/z/my-project/download/MyZipVault-Skills-Checklist-Import-Template.xlsx"

# ── Starter content: built-in VaultVerify checklists mapped to Zipvault schema ──
MEDSURG = [  # (category, skill, high-risk)
    ("IV Therapy & Access", "IV insertion & site care", False),
    ("Medication Administration", "IV push / infusion medications", False),
    ("Medication Administration", "Medication administration (oral / IM / SubQ)", False),
    ("Medication Administration", "Diabetic care & insulin management", True),
    ("Wound Care & Skin", "Wound care & dressing changes", False),
    ("Symptom & Pain Management", "Pain management & reassessment", False),
    ("Genitourinary Care", "Foley catheter insertion & care", False),
    ("Respiratory Care", "Tracheostomy care & suctioning", False),
    ("Respiratory Care", "Oxygen therapy & respiratory treatments", False),
    ("Care Planning & Education", "Care plans & discharge education", False),
    ("Safety & Compliance", "Fall risk & safety protocols", False),
    ("Cardiac Monitoring", "Telemetry — basic rhythm interpretation", False),
    ("Transfusion Medicine", "Blood transfusion administration", True),
    ("GI & Nutrition", "NG / G-tube feeding & care", False),
]
ICU = [
    ("Respiratory Support", "Ventilator management", True),
    ("Hemodynamic Monitoring", "Titrating vasoactive drips (norepi, vasopressin)", True),
    ("Hemodynamic Monitoring", "Arterial lines & hemodynamic monitoring", True),
    ("Vascular Access & Infection Control", "Central line care & CLABSI prevention", True),
    ("Renal Replacement Therapy", "CRRT / dialysis circuits", True),
    ("Chest Tubes & Drains", "Chest tube management", True),
    ("Neurological Care", "ICP monitoring & neuro assessments", False),
    ("Sedation & Analgesia", "Sedation & analgesia (RASS scoring)", False),
    ("Emergency Response", "Code Blue / ACLS response", False),
    ("Transfusion Medicine", "Blood transfusion in critical care", True),
    ("Airway & Tracheostomy Care", "Tracheostomy care & suctioning", False),
    ("Wound Care & Skin", "Pressure injury prevention", False),
    ("Psychosocial & End-of-Life", "Family communication & end-of-life care", False),
]
SPEC_EXAMPLES = [  # the 5 format examples from the Zipvault spec, verbatim
    ("Medication Administration", "Documentation on M.A.R."),
    ("Medication Administration", "Dose Calculation"),
    ("Medication Administration", "Generic Equivalents"),
    ("Medication Administration", "Usage of PDR"),
    ("Medication Administration", "Knowledge of Drug Actions/Interactions"),
]

wb = Workbook()

# ════════════════════════════════ Sheet 1: Instructions ══════════════════════
ws = wb.active
ws.title = "Instructions"
setup_sheet(ws, title="MyZipVault Skills Checklist — Import Template", last_col=6)

r = 4
def line(text, style="body", indent=0):
    global r
    c = ws.cell(row=r, column=2 + indent, value=text)
    c.font = {"body": font_body(), "sub": font_subheader(), "cap": font_caption()}[style]
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
    ws.row_dimensions[r].height = 20 if style != "cap" else 16
    r += 1

line("Use this workbook to load skill checklists into VaultVerify (Zipvault-compatible format v1).")
line("Upload path: VaultVerify → Super Admin → Skills → Import workbook (.xlsx).")
r += 1

line("Rules", "sub")
line("1.  Do not change the column headers in the Skills Data sheet.")
line("2.  Profession must be one of: Nursing, Allied, Pharmacy, Locums.")
line("3.  Question Type must be one of: rating_1_4, yes_no, text.")
line("4.  Has N/A Option must be: Yes or No.")
line("5.  Do not leave Profession, Job Title, Specialty, Category, or Skill Name blank.")
line("6.  Save the file as .xlsx before uploading.")
line("7.  Re-importing the same Profession + Job Title + Specialty + Skill Name updates that row (safe to re-upload).")
r += 1

line("Rating scale (rating_1_4)", "sub")
line("1  =  No theory and/or experience  →  VaultVerify: Never performed")
line("2  =  Limited Experience  →  VaultVerify: With supervision")
line("3  =  Experienced / minimal support needed  →  VaultVerify: Independent")
line("4  =  Proficient  →  VaultVerify: Can teach others")
r += 1

line("Question types", "sub")
line("rating_1_4  —  reference or nurse rates the skill on the 4-point scale above.")
line("yes_no  —  confirmatory check: the skill is (or is not) signed off.")
line("text  —  free-text attestation, e.g. certification numbers or comments.")
r += 1

line("Format example rows (Specialty 'General')", "sub")
hdr = ["Profession", "Job Title", "Specialty", "Category", "Skill Name", "Question Type", "Has N/A Option"]
ex_row = r
for col, h in enumerate(hdr, 2):
    ws.cell(row=ex_row, column=col, value=h)
style_header_row(ws, ex_row, 2, 8)
r += 1
for i, (cat, skill) in enumerate(SPEC_EXAMPLES):
    vals = ["Nursing", "RN", "General", cat, skill, "rating_1_4", "No"]
    for col, v in enumerate(vals, 2):
        ws.cell(row=r, column=col, value=v)
    style_data_row(ws, r, 2, 8, i)
    r += 1
r += 1
line("The Skills Data sheet ships with these example rows plus full Med-Surg and ICU / Critical Care", "cap")
line("starter content for RN. Edit or replace rows freely — the importer upserts and never duplicates.", "cap")

auto_fit_columns(ws, header_row=ex_row, data_start_row=ex_row + 1)
ws.column_dimensions["B"].width = 13
for col in "CDEFGH":
    ws.column_dimensions[col].width = 22

# ════════════════════════════════ Sheet 2: Skills Data ═══════════════════════
sd = wb.create_sheet("Skills Data")
setup_sheet(sd, title="Skills Data", last_col=8)
HDR_ROW = 4
for col, h in enumerate(hdr, 2):
    sd.cell(row=HDR_ROW, column=col, value=h)
style_header_row(sd, HDR_ROW, 2, 8)

rows = []
for cat, skill, hr in MEDSURG:
    rows.append(("Nursing", "RN", "Med-Surg", cat, skill, "rating_1_4", "Yes"))
for cat, skill, hr in ICU:
    rows.append(("Nursing", "RN", "ICU / Critical Care", cat, skill, "rating_1_4", "Yes"))
for cat, skill in SPEC_EXAMPLES:
    rows.append(("Nursing", "RN", "General", cat, skill, "rating_1_4", "No"))

rr = HDR_ROW + 1
for i, vals in enumerate(rows):
    for col, v in enumerate(vals, 2):
        sd.cell(row=rr, column=col, value=v)
    style_data_row(sd, rr, 2, 8, i)
    rr += 1

last = rr + 300  # room for user additions
dv_prof = DataValidation(type="list", formula1='"Nursing,Allied,Pharmacy,Locums"', allow_blank=True, showErrorMessage=True)
dv_type = DataValidation(type="list", formula1='"rating_1_4,yes_no,text"', allow_blank=True, showErrorMessage=True)
dv_na   = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True, showErrorMessage=True)
dv_prof.error = "Profession must be: Nursing, Allied, Pharmacy, or Locums"
dv_type.error = "Question Type must be: rating_1_4, yes_no, or text"
dv_na.error = "Has N/A Option must be: Yes or No"
sd.add_data_validation(dv_prof); sd.add_data_validation(dv_type); sd.add_data_validation(dv_na)
dv_prof.add(f"B{HDR_ROW+1}:B{last}")
dv_type.add(f"G{HDR_ROW+1}:G{last}")
dv_na.add(f"H{HDR_ROW+1}:H{last}")

note = sd.cell(row=rr + 1, column=2, value="Rows above ship with VaultVerify Med-Surg + ICU starter content for RN. Add your own rows below — dropdowns stay active through row %d." % last)
note.font = font_caption()
sd.freeze_panes = "B5"
auto_fit_columns(sd, header_row=HDR_ROW, data_start_row=HDR_ROW + 1)
sd.column_dimensions["F"].width = 30   # Skill Name
sd.column_dimensions["E"].width = 28   # Category
sd.column_dimensions["D"].width = 18   # Specialty

wb.properties.creator = "Z.ai"
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print("saved", OUT, "| data rows:", len(rows))
